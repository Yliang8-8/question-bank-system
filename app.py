from functools import wraps
from html import escape
from pathlib import Path
import csv
import io
import math
import re
import sqlite3

import openpyxl
from flask import Flask, flash, g, redirect, render_template, request, send_file, session, url_for
from werkzeug.security import check_password_hash, generate_password_hash


BASE_DIR = Path(__file__).resolve().parent
DATABASE = BASE_DIR / "question_bank.db"
EXPORT_DIR = BASE_DIR / "exports"

app = Flask(__name__)
app.config["SECRET_KEY"] = "question-bank-dev-key"

QUESTION_TYPE_LABELS = {
    "single": "单选题",
    "multiple": "多选题",
    "true_false": "判断题",
    "fill": "填空题",
}
QUESTION_TYPE_ORDER = ("single", "multiple", "true_false", "fill")
QUESTION_TYPE_RANK = {question_type: index for index, question_type in enumerate(QUESTION_TYPE_ORDER)}

DIFFICULTY_LABELS = {
    "easy": "简单",
    "medium": "中等",
    "hard": "困难",
}
DIFFICULTY_RANK = {difficulty: index for index, difficulty in enumerate(DIFFICULTY_LABELS)}

SUBJECTS = ["计算机基础", "高等数学", "英语"]
DEFAULT_SUBJECT = "计算机基础"

DEFAULT_QUESTIONS = [
    ("Python 中用于定义函数的关键字是什么？", "single", "easy", "Python基础", "A", "def 用于定义函数。", "A. def\nB. class\nC. import\nD. return"),
    ("下列哪些属于关系型数据库？", "multiple", "medium", "数据库基础", "A,C", "MySQL 和 SQLite 都是关系型数据库。", "A. MySQL\nB. Redis\nC. SQLite\nD. MongoDB"),
    ("HTTP 是一种应用层协议。", "true_false", "easy", "计算机网络", "正确", "HTTP 工作在应用层。", ""),
    ("Flask 属于什么类型的 Web 框架？", "single", "medium", "Web开发", "A", "Flask 是轻量级 Web 框架。", "A. 轻量级 Web 框架\nB. 操作系统\nC. 数据库系统\nD. 图像处理工具"),
    ("SQLite 适合做原型开发和小型项目。", "true_false", "easy", "数据库基础", "正确", "SQLite 轻量、免安装，适合演示和原型。", ""),
    ("Python 中哪个符号用于单行注释？", "single", "easy", "Python基础", "A", "Python 使用 # 开始单行注释。", "A. #\nB. //\nC. <!-- -->\nD. /* */"),
    ("下列哪些是 Python 的基本数据类型？", "multiple", "easy", "Python基础", "A,B,C", "整数、字符串和列表都是常用基本数据类型。", "A. int\nB. str\nC. list\nD. table"),
    ("列表是一种有序且可变的数据结构。", "true_false", "easy", "Python基础", "正确", "Python 列表支持增删改操作。", ""),
    ("SQL 查询数据通常使用哪个关键字？", "single", "easy", "数据库基础", "A", "SELECT 用于查询数据。", "A. SELECT\nB. DELETE\nC. UPDATE\nD. INSERT"),
    ("下列哪些属于 SQL 的数据定义语言操作？", "multiple", "medium", "数据库基础", "A,B", "CREATE 和 ALTER 用于定义或修改数据库对象。", "A. CREATE\nB. ALTER\nC. SELECT\nD. INSERT"),
    ("主键字段的值可以重复。", "true_false", "easy", "数据库基础", "错误", "主键用于唯一标识记录，不能重复。", ""),
    ("TCP 具有什么特点？", "multiple", "medium", "计算机网络", "A,B", "TCP 提供面向连接和可靠的数据传输。", "A. 面向连接\nB. 可靠传输\nC. 不校验数据\nD. 只用于静态网页"),
    ("IP 地址用于标识网络中的设备。", "true_false", "easy", "计算机网络", "正确", "IP 地址用于网络寻址。", ""),
    ("HTML 主要负责网页的什么部分？", "single", "easy", "Web开发", "A", "HTML 负责网页结构。", "A. 页面结构\nB. 数据备份\nC. 服务器调度\nD. 图像压缩"),
    ("CSS 主要用于描述网页的结构。", "true_false", "easy", "Web开发", "错误", "CSS 主要负责网页样式和布局。", ""),
    ("下列哪些属于前端开发技术？", "multiple", "medium", "Web开发", "A,B,C", "HTML、CSS 和 JavaScript 都常用于前端开发。", "A. HTML\nB. CSS\nC. JavaScript\nD. SQLite"),
    ("REST API 通常使用 HTTP 方法表达资源操作。", "true_false", "medium", "Web开发", "正确", "GET、POST、PUT、DELETE 是常见 HTTP 方法。", ""),
    ("算法的时间复杂度用于描述什么？", "single", "medium", "算法基础", "A", "时间复杂度描述输入规模增长时运行时间的变化趋势。", "A. 运行时间随输入规模增长的趋势\nB. 程序界面颜色\nC. 数据库容量上限\nD. 网络带宽大小"),
    ("二分查找要求数据通常具有什么特征？", "single", "medium", "算法基础", "A", "二分查找通常要求数据有序。", "A. 有序\nB. 完全随机\nC. 只能是字符串\nD. 必须为空"),
    ("栈遵循先进先出的原则。", "true_false", "easy", "数据结构", "错误", "栈遵循后进先出，队列才是先进先出。", ""),
    ("下列哪些属于常见排序算法？", "multiple", "easy", "算法基础", "A,B,C", "冒泡排序、快速排序和归并排序都是排序算法。", "A. 冒泡排序\nB. 快速排序\nC. 归并排序\nD. 页面渲染"),
    ("面向对象程序设计常见的基本特征包括哪些？", "multiple", "medium", "程序设计", "A,B,C", "封装、继承和多态是常见特征。", "A. 封装\nB. 继承\nC. 多态\nD. 网络寻址"),
    ("版本控制工具 Git 可以记录代码变更历史。", "true_false", "easy", "软件工程", "正确", "Git 用于版本管理和协作开发。", ""),
    ("软件测试的主要目的是什么？", "single", "medium", "软件工程", "A", "软件测试用于发现问题并验证系统质量。", "A. 发现问题并验证质量\nB. 删除全部代码\nC. 替代需求分析\nD. 改变硬件结构"),
    ("Python 中用于输出内容的常用函数是____。", "fill", "easy", "Python基础", "print", "print 函数常用于在控制台输出内容。", ""),
    ("数据库中用于唯一标识一条记录的字段通常称为____。", "fill", "medium", "数据库基础", "主键", "主键用于唯一标识表中的记录。", ""),
    ("网页中负责交互逻辑的脚本语言通常是____。", "fill", "easy", "Web开发", "JavaScript", "JavaScript 常用于实现网页交互逻辑。", ""),
    ("软件开发中用于保存代码变更历史的工具类别是____。", "fill", "medium", "软件工程", "版本控制", "版本控制工具用于记录和管理代码变更。", ""),
]


def with_subject(subject, rows):
    return [
        (content, question_type, difficulty, subject, chapter, answer, analysis, options)
        for content, question_type, difficulty, chapter, answer, analysis, options in rows
    ]


DEFAULT_QUESTIONS += with_subject(
    "高等数学",
    [
        ("函数 y=x^2 在 x=2 处的导数是多少？", "single", "easy", "导数与微分", "B", "y'=2x，所以 x=2 时导数为 4。", "A. 2\nB. 4\nC. 6\nD. 8"),
        ("lim(x→0) sin x / x 的值是多少？", "single", "easy", "极限与连续", "A", "这是第一重要极限，值为 1。", "A. 1\nB. 0\nC. -1\nD. 不存在"),
        ("函数 y=e^x 的导数是____。", "fill", "easy", "导数与微分", "e^x", "指数函数 e^x 的导数仍为 e^x。", ""),
        ("若函数在某点可导，则它在该点一定连续。", "true_false", "easy", "导数与微分", "正确", "可导必连续，但连续不一定可导。", ""),
        ("下列哪些函数是偶函数？", "multiple", "medium", "函数性质", "A,C", "cos x 与 x^2 都满足 f(-x)=f(x)。", "A. cos x\nB. sin x\nC. x^2\nD. x^3"),
        ("∫ 2x dx 的一个原函数是____。", "fill", "easy", "不定积分", "x^2", "x^2 的导数为 2x。", ""),
        ("函数 y=ln x 的定义域是？", "single", "easy", "函数性质", "C", "ln x 只在 x>0 时有定义。", "A. 全体实数\nB. x≠0\nC. x>0\nD. x<0"),
        ("定积分表示曲边梯形面积时，结果一定为正数。", "true_false", "medium", "定积分", "错误", "定积分是代数量，函数在 x 轴下方时贡献为负。", ""),
        ("函数 y=sin x 的导数是？", "single", "easy", "导数与微分", "A", "sin x 的导数为 cos x。", "A. cos x\nB. -cos x\nC. tan x\nD. -sin x"),
        ("微分中值定理要求函数在闭区间连续、开区间可导。", "true_false", "medium", "中值定理", "正确", "这是拉格朗日中值定理的基本条件。", ""),
        ("下列哪些属于不定积分的基本性质？", "multiple", "medium", "不定积分", "A,B", "线性性质包括常数倍可提出、和差可拆分。", "A. 常数倍可提出\nB. 和差可拆分\nC. 积分上下限固定\nD. 结果唯一没有常数"),
        ("lim(x→∞) 1/x 的值是多少？", "single", "easy", "极限与连续", "B", "当 x 趋于无穷大时，1/x 趋于 0。", "A. 1\nB. 0\nC. 无穷大\nD. 不存在"),
        ("函数 y=x^3 在 x=1 处的导数是____。", "fill", "medium", "导数与微分", "3", "y'=3x^2，x=1 时为 3。", ""),
        ("连续函数在闭区间上一定有最大值和最小值。", "true_false", "medium", "极限与连续", "正确", "闭区间上连续函数满足最大最小值定理。", ""),
        ("下列哪些是无穷小量？", "multiple", "medium", "极限与连续", "A,C", "x→0 时 x 和 x^2 都趋于 0。", "A. x（x→0）\nB. 1/x（x→0）\nC. x^2（x→0）\nD. 1"),
        ("∫ 1/x dx 的结果是？", "single", "medium", "不定积分", "B", "1/x 的原函数是 ln|x|+C。", "A. x\nB. ln|x|+C\nC. e^x+C\nD. x^2+C"),
        ("矩阵的行列式只对方阵定义。", "true_false", "easy", "线性代数基础", "正确", "行列式是方阵的重要数值特征。", ""),
        ("二阶行列式 |a b; c d| 的值是____。", "fill", "medium", "线性代数基础", "ad-bc", "二阶行列式计算公式为 ad-bc。", ""),
        ("下列哪些属于求极限常用方法？", "multiple", "medium", "极限与连续", "A,B,C", "因式分解、有理化、洛必达法则都常用于求极限。", "A. 因式分解\nB. 有理化\nC. 洛必达法则\nD. 改变坐标轴颜色"),
        ("函数 y=1/x 在 x=0 处连续。", "true_false", "easy", "极限与连续", "错误", "该函数在 x=0 处没有定义，因此不连续。", ""),
        ("函数 y=cos x 的导数是？", "single", "easy", "导数与微分", "D", "cos x 的导数是 -sin x。", "A. sin x\nB. cos x\nC. -cos x\nD. -sin x"),
        ("若二元函数 z=f(x,y)，则偏导数表示函数沿某一变量方向的变化率。", "true_false", "medium", "多元函数微分", "正确", "偏导数刻画保持其他变量不变时的变化率。", ""),
        ("∫ cos x dx 的一个原函数是____。", "fill", "easy", "不定积分", "sin x", "sin x 的导数为 cos x。", ""),
        ("下列哪些说法正确？", "multiple", "hard", "级数", "A,D", "收敛级数有确定和；正项级数可用比较判别法。", "A. 收敛级数有确定和\nB. 发散级数也一定有有限和\nC. 所有级数都收敛\nD. 正项级数可用比较判别法"),
        ("函数极值点处导数一定为 0。", "true_false", "medium", "导数应用", "错误", "可导极值点导数为 0，不可导极值点不满足。", ""),
        ("一阶线性微分方程的未知量是？", "single", "medium", "微分方程", "C", "微分方程要求解的是未知函数。", "A. 常数\nB. 矩阵\nC. 函数\nD. 表格"),
        ("定积分 ∫_0^1 x dx 的值是____。", "fill", "easy", "定积分", "1/2", "x 的原函数是 x^2/2，代入 0 到 1 得 1/2。", ""),
        ("下列哪些属于导数应用？", "multiple", "medium", "导数应用", "A,B,C", "导数可用于判断单调性、极值和曲线切线斜率。", "A. 判断单调性\nB. 求极值\nC. 求切线斜率\nD. 设置网页字体"),
        ("如果两个向量线性无关，则其中一个一定可以由另一个线性表示。", "true_false", "medium", "线性代数基础", "错误", "线性无关意味着不能互相线性表示。", ""),
        ("函数 y=x^4 的二阶导数是____。", "fill", "medium", "导数与微分", "12x^2", "一阶导数为 4x^3，二阶导数为 12x^2。", ""),
    ],
)

DEFAULT_QUESTIONS += with_subject(
    "英语",
    [
        ("Choose the correct word: She ____ to school every day.", "single", "easy", "语法", "B", "第三人称单数主语 She 后动词用 goes。", "A. go\nB. goes\nC. going\nD. gone"),
        ("Which words are nouns?", "multiple", "easy", "词汇", "A,C", "book 和 teacher 是名词。", "A. book\nB. quickly\nC. teacher\nD. beautiful"),
        ("The plural form of 'child' is ____.", "fill", "easy", "词汇", "children", "child 的复数形式是不规则变化 children。", ""),
        ("'He is reading a book' is present continuous tense.", "true_false", "easy", "语法", "正确", "be doing 表示现在进行时。", ""),
        ("Choose the correct article: I saw ____ apple on the table.", "single", "easy", "语法", "A", "apple 以元音音素开头，前面用 an。", "A. an\nB. a\nC. the\nD. /"),
        ("Which words are adjectives?", "multiple", "medium", "词汇", "B,D", "beautiful 和 careful 是形容词。", "A. run\nB. beautiful\nC. slowly\nD. careful"),
        ("The opposite of 'hot' is ____.", "fill", "easy", "词汇", "cold", "hot 的反义词是 cold。", ""),
        ("'They was late' is grammatically correct.", "true_false", "easy", "语法", "错误", "They 后应使用 were。", ""),
        ("Choose the correct form: We ____ English now.", "single", "easy", "语法", "C", "now 常提示现在进行时，主语 We 用 are studying。", "A. study\nB. studied\nC. are studying\nD. studies"),
        ("Which are modal verbs?", "multiple", "medium", "语法", "A,B,D", "can、must、should 都是情态动词。", "A. can\nB. must\nC. desk\nD. should"),
        ("The past tense of 'write' is ____.", "fill", "easy", "词汇", "wrote", "write 的过去式是 wrote。", ""),
        ("A topic sentence usually states the main idea of a paragraph.", "true_false", "medium", "阅读理解", "正确", "主题句通常概括段落中心。", ""),
        ("Choose the correct translation of 'environment'.", "single", "easy", "词汇", "D", "environment 表示环境。", "A. 经济\nB. 教育\nC. 能源\nD. 环境"),
        ("Which phrases can express cause and effect?", "multiple", "medium", "写作", "A,C", "because of 和 therefore 可表达因果。", "A. because of\nB. in contrast\nC. therefore\nD. on the left"),
        ("The comparative form of 'good' is ____.", "fill", "easy", "词汇", "better", "good 的比较级是不规则形式 better。", ""),
        ("'Although' is commonly used to introduce contrast.", "true_false", "medium", "写作", "正确", "although 常引导让步或转折关系。", ""),
        ("Choose the correct sentence.", "single", "medium", "语法", "C", "第三个句子主谓和时态搭配正确。", "A. He have a book.\nB. She go home.\nC. They are students.\nD. I is happy."),
        ("Which words are verbs?", "multiple", "easy", "词汇", "A,D", "read 和 improve 是动词。", "A. read\nB. desk\nC. blue\nD. improve"),
        ("The synonym of 'important' is ____.", "fill", "medium", "词汇", "significant", "significant 可表示重要的、意义重大的。", ""),
        ("In English writing, every paragraph must have exactly three sentences.", "true_false", "medium", "写作", "错误", "段落句数没有固定必须为三句的要求。", ""),
        ("Choose the correct preposition: interested ____ science.", "single", "easy", "语法", "B", "固定搭配为 be interested in。", "A. at\nB. in\nC. on\nD. by"),
        ("Which expressions are useful in academic writing?", "multiple", "hard", "写作", "A,B,C", "in addition、for example、as a result 都常用于学术写作。", "A. in addition\nB. for example\nC. as a result\nD. gonna"),
        ("The past participle of 'see' is ____.", "fill", "medium", "词汇", "seen", "see 的过去分词是 seen。", ""),
        ("Scanning means reading quickly to find specific information.", "true_false", "medium", "阅读理解", "正确", "scanning 是寻读，用于快速定位特定信息。", ""),
        ("Choose the correct passive voice: The bridge ____ last year.", "single", "medium", "语法", "A", "被动语态结构为 was/were + 过去分词。", "A. was built\nB. built\nC. builds\nD. is build"),
        ("Which sentences are questions?", "multiple", "easy", "语法", "A,C", "带疑问结构或问号的句子是疑问句。", "A. Are you ready?\nB. She is ready.\nC. What is your name?\nD. He likes math."),
        ("The word 'efficient' means ____ in Chinese.", "fill", "medium", "词汇", "高效的", "efficient 表示高效的。", ""),
        ("Skimming means reading for the general idea.", "true_false", "medium", "阅读理解", "正确", "skimming 是略读，用于获取文章大意。", ""),
        ("Choose the correct connector: It was raining; ____, we stayed at home.", "single", "medium", "写作", "B", "therefore 表示因此，符合因果关系。", "A. however\nB. therefore\nC. although\nD. unless"),
        ("Which words are related to engineering English?", "multiple", "medium", "词汇", "A,B,D", "system、design、process 都常见于工程英语。", "A. system\nB. design\nC. orange\nD. process"),
    ],
)


def question_subject(question):
    return question[3] if len(question) == 8 else DEFAULT_SUBJECT


def default_question_row(question):
    if len(question) == 8:
        return question
    content, question_type, difficulty, chapter, answer, analysis, options = question
    return (content, question_type, difficulty, DEFAULT_SUBJECT, chapter, answer, analysis, options)


DEFAULT_OPTION_MAP = {question[0]: question[-1] for question in DEFAULT_QUESTIONS}
DEFAULT_SUBJECT_MAP = {question[0]: question_subject(question) for question in DEFAULT_QUESTIONS}


def get_db():
    if "db" not in g:
        g.db = sqlite3.connect(DATABASE)
        g.db.row_factory = sqlite3.Row
        g.db.execute("PRAGMA foreign_keys = ON")
    return g.db


@app.teardown_appcontext
def close_db(_error=None):
    db = g.pop("db", None)
    if db is not None:
        db.close()


def init_db():
    db = get_db()
    db.executescript(
        """
        CREATE TABLE IF NOT EXISTS users (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            username TEXT NOT NULL UNIQUE,
            display_name TEXT NOT NULL DEFAULT '',
            password_hash TEXT NOT NULL,
            role TEXT NOT NULL CHECK(role IN ('student', 'teacher'))
        );

        CREATE TABLE IF NOT EXISTS classes (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL,
            teacher_id INTEGER NOT NULL,
            created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (teacher_id) REFERENCES users(id)
        );

        CREATE TABLE IF NOT EXISTS questions (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            content TEXT NOT NULL,
            question_type TEXT NOT NULL CHECK(question_type IN ('single', 'multiple', 'true_false', 'fill')),
            difficulty TEXT NOT NULL CHECK(difficulty IN ('easy', 'medium', 'hard')),
            subject TEXT NOT NULL DEFAULT '计算机基础',
            chapter TEXT NOT NULL DEFAULT '',
            answer TEXT NOT NULL,
            analysis TEXT NOT NULL DEFAULT '',
            options TEXT NOT NULL DEFAULT '',
            created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
        );

        CREATE TABLE IF NOT EXISTS papers (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            title TEXT NOT NULL,
            total_score INTEGER NOT NULL,
            duration_minutes INTEGER NOT NULL DEFAULT 30,
            max_attempts INTEGER NOT NULL DEFAULT 1,
            created_by INTEGER NOT NULL,
            created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (created_by) REFERENCES users(id)
        );

        CREATE TABLE IF NOT EXISTS paper_classes (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            paper_id INTEGER NOT NULL,
            class_id INTEGER NOT NULL,
            FOREIGN KEY (paper_id) REFERENCES papers(id) ON DELETE CASCADE,
            FOREIGN KEY (class_id) REFERENCES classes(id) ON DELETE CASCADE,
            UNIQUE(paper_id, class_id)
        );

        CREATE TABLE IF NOT EXISTS paper_questions (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            paper_id INTEGER NOT NULL,
            question_id INTEGER NOT NULL,
            score INTEGER NOT NULL,
            sort_order INTEGER NOT NULL,
            FOREIGN KEY (paper_id) REFERENCES papers(id) ON DELETE CASCADE,
            FOREIGN KEY (question_id) REFERENCES questions(id)
        );

        CREATE TABLE IF NOT EXISTS exam_attempts (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            paper_id INTEGER NOT NULL,
            user_id INTEGER NOT NULL,
            status TEXT NOT NULL DEFAULT 'submitted',
            score INTEGER NOT NULL,
            total_score INTEGER NOT NULL,
            started_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
            submitted_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (paper_id) REFERENCES papers(id),
            FOREIGN KEY (user_id) REFERENCES users(id)
        );

        CREATE TABLE IF NOT EXISTS exam_answers (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            attempt_id INTEGER NOT NULL,
            paper_question_id INTEGER NOT NULL,
            answer TEXT NOT NULL DEFAULT '',
            is_correct INTEGER NOT NULL DEFAULT 0,
            score INTEGER NOT NULL DEFAULT 0,
            FOREIGN KEY (attempt_id) REFERENCES exam_attempts(id) ON DELETE CASCADE,
            FOREIGN KEY (paper_question_id) REFERENCES paper_questions(id)
        );
        """
    )

    question_columns = {
        row["name"] for row in db.execute("PRAGMA table_info(questions)").fetchall()
    }
    if "options" not in question_columns:
        db.execute("ALTER TABLE questions ADD COLUMN options TEXT NOT NULL DEFAULT ''")
    if "subject" not in question_columns:
        db.execute("ALTER TABLE questions ADD COLUMN subject TEXT NOT NULL DEFAULT '计算机基础'")
    if "source" not in question_columns:
        db.execute("ALTER TABLE questions ADD COLUMN source TEXT NOT NULL DEFAULT 'manual'")
    db.execute(
        """
        UPDATE questions
        SET source = 'import'
        WHERE source = 'manual' AND content LIKE '【批量验证】%'
        """
    )

    user_columns = {
        row["name"] for row in db.execute("PRAGMA table_info(users)").fetchall()
    }
    if "display_name" not in user_columns:
        db.execute("ALTER TABLE users ADD COLUMN display_name TEXT NOT NULL DEFAULT ''")
    if "class_id" not in user_columns:
        db.execute("ALTER TABLE users ADD COLUMN class_id INTEGER REFERENCES classes(id)")

    paper_columns = {
        row["name"] for row in db.execute("PRAGMA table_info(papers)").fetchall()
    }
    if "duration_minutes" not in paper_columns:
        db.execute("ALTER TABLE papers ADD COLUMN duration_minutes INTEGER NOT NULL DEFAULT 30")
    if "max_attempts" not in paper_columns:
        db.execute("ALTER TABLE papers ADD COLUMN max_attempts INTEGER NOT NULL DEFAULT 1")

    if db.execute("SELECT COUNT(*) AS count FROM users").fetchone()["count"] == 0:
        db.executemany(
            "INSERT INTO users (username, display_name, password_hash, role) VALUES (?, ?, ?, ?)",
            [
                ("teacher", "教师测试账号", generate_password_hash("teacher123"), "teacher"),
                ("student", "学生测试账号", generate_password_hash("student123"), "student"),
            ],
        )
    db.execute(
        """
        UPDATE users
        SET display_name = CASE username
            WHEN 'teacher' THEN '教师测试账号'
            WHEN 'student' THEN '学生测试账号'
            ELSE username
        END
        WHERE TRIM(COALESCE(display_name, '')) = ''
        """
    )

    existing_questions = {
        row["content"] for row in db.execute("SELECT content FROM questions").fetchall()
    }
    missing_questions = [
        question for question in DEFAULT_QUESTIONS if question[0] not in existing_questions
    ]
    if missing_questions:
        db.executemany(
            """
            INSERT INTO questions
            (content, question_type, difficulty, subject, chapter, answer, analysis, options)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?)
            """,
            [default_question_row(question) for question in missing_questions],
        )
    for content, subject in DEFAULT_SUBJECT_MAP.items():
        db.execute(
            """
            UPDATE questions SET subject = ?
            WHERE content = ? AND TRIM(COALESCE(subject, '')) = ''
            """,
            (subject, content),
        )
    for content, options in DEFAULT_OPTION_MAP.items():
        db.execute(
            """
            UPDATE questions SET options = ?
            WHERE content = ? AND TRIM(COALESCE(options, '')) = ''
            """,
            (options, content),
        )
    db.commit()


@app.before_request
def prepare_database():
    init_db()


@app.context_processor
def inject_globals():
    return {
        "current_user": session.get("user"),
        "question_type_labels": QUESTION_TYPE_LABELS,
        "difficulty_labels": DIFFICULTY_LABELS,
        "subjects": SUBJECTS,
        "parse_options": parse_options,
    }


def login_required(view):
    @wraps(view)
    def wrapped_view(**kwargs):
        if "user" not in session:
            return redirect(url_for("login"))
        return view(**kwargs)

    return wrapped_view


def teacher_required(view):
    @wraps(view)
    def wrapped_view(**kwargs):
        if session.get("user", {}).get("role") != "teacher":
            flash("只有教师账号可以执行此操作。", "error")
            return redirect(url_for("dashboard"))
        return view(**kwargs)

    return wrapped_view


@app.route("/")
def index():
    if "user" not in session:
        return redirect(url_for("login"))
    return redirect(url_for("dashboard"))


@app.route("/login", methods=("GET", "POST"))
def login():
    if "user" in session:
        return redirect(url_for("dashboard"))
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "")
        user = get_db().execute("SELECT * FROM users WHERE username = ?", (username,)).fetchone()

        if user is None or not check_password_hash(user["password_hash"], password):
            flash("账号或密码错误。", "error")
        else:
            session.clear()
            session["user"] = {
                "id": user["id"],
                "username": user["username"],
                "display_name": user["display_name"] or user["username"],
                "role": user["role"],
            }
            return redirect(url_for("dashboard"))
    return render_template("login.html")


@app.route("/register", methods=("GET", "POST"))
def register():
    if "user" in session:
        return redirect(url_for("dashboard"))

    form = {"username": "", "display_name": "", "role": "student"}
    if request.method == "POST":
        form = {
            "username": request.form.get("username", "").strip(),
            "display_name": request.form.get("display_name", "").strip(),
            "role": request.form.get("role", "student").strip(),
        }
        error = validate_register_form(form)
        if error:
            flash(error, "error")
        else:
            initial_password = f"czu{form['username']}"
            try:
                get_db().execute(
                    "INSERT INTO users (username, display_name, password_hash, role) VALUES (?, ?, ?, ?)",
                    (
                        form["username"],
                        form["display_name"],
                        generate_password_hash(initial_password),
                        form["role"],
                    ),
                )
                get_db().commit()
            except sqlite3.IntegrityError:
                flash("该账号已存在，请换一个 8 位数字账号。", "error")
            else:
                flash(f"注册成功，初始密码为 {initial_password}，登录后请及时修改。", "success")
                return redirect(url_for("login"))
    return render_template("register.html", form=form)


@app.get("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))


@app.route("/password", methods=("GET", "POST"))
@login_required
def change_password():
    user = get_db().execute(
        "SELECT * FROM users WHERE id = ?",
        (session["user"]["id"],),
    ).fetchone()
    if user is None:
        session.clear()
        flash("账号状态异常，请重新登录。", "error")
        return redirect(url_for("login"))

    if request.method == "POST":
        display_name = request.form.get("display_name", "").strip()
        old_password = request.form.get("old_password", "")
        new_password = request.form.get("new_password", "")
        confirm_password = request.form.get("confirm_password", "")
        password_changed = any([old_password, new_password, confirm_password])

        if not valid_display_name(display_name):
            flash("用户名需为 2-20 个字符。", "error")
        elif password_changed and not check_password_hash(user["password_hash"], old_password):
            flash("原密码不正确。", "error")
        elif password_changed and len(new_password) < 6:
            flash("新密码至少需要 6 位。", "error")
        elif password_changed and new_password != confirm_password:
            flash("两次输入的新密码不一致。", "error")
        else:
            if password_changed:
                get_db().execute(
                    "UPDATE users SET display_name = ?, password_hash = ? WHERE id = ?",
                    (display_name, generate_password_hash(new_password), session["user"]["id"]),
                )
            else:
                get_db().execute(
                    "UPDATE users SET display_name = ? WHERE id = ?",
                    (display_name, session["user"]["id"]),
                )
            get_db().commit()
            session["user"]["display_name"] = display_name
            session.modified = True
            if password_changed:
                session.clear()
                flash("账号设置已保存，密码已修改，请使用新密码重新登录。", "success")
                return redirect(url_for("login"))
            flash("账号设置已保存。", "success")
            return redirect(url_for("dashboard"))
    return render_template("change_password.html", user=user)


@app.get("/dashboard")
@login_required
def dashboard():
    db = get_db()
    stats = {
        "questions": db.execute("SELECT COUNT(*) AS count FROM questions").fetchone()["count"],
        "easy": db.execute("SELECT COUNT(*) AS count FROM questions WHERE difficulty = 'easy'").fetchone()["count"],
        "medium": db.execute("SELECT COUNT(*) AS count FROM questions WHERE difficulty = 'medium'").fetchone()["count"],
        "hard": db.execute("SELECT COUNT(*) AS count FROM questions WHERE difficulty = 'hard'").fetchone()["count"],
        "papers": db.execute("SELECT COUNT(*) AS count FROM papers").fetchone()["count"],
        "attempts": db.execute("SELECT COUNT(*) AS count FROM exam_attempts").fetchone()["count"],
    }
    new_counts = db.execute(
        """
        SELECT COUNT(*) AS questions,
               COALESCE(SUM(CASE WHEN difficulty = 'easy' THEN 1 ELSE 0 END), 0) AS easy,
               COALESCE(SUM(CASE WHEN difficulty = 'medium' THEN 1 ELSE 0 END), 0) AS medium,
               COALESCE(SUM(CASE WHEN difficulty = 'hard' THEN 1 ELSE 0 END), 0) AS hard
        FROM questions
        WHERE source = 'import'
        """
    ).fetchone()
    student_class = None
    if session["user"]["role"] == "student":
        student_class = db.execute(
            """
            SELECT c.name
            FROM users u
            LEFT JOIN classes c ON c.id = u.class_id
            WHERE u.id = ?
            """,
            (session["user"]["id"],),
        ).fetchone()
    recent_papers = db.execute(
        """
        SELECT p.id, p.title, p.total_score, p.created_at,
               COALESCE(NULLIF(u.display_name, ''), u.username) AS creator
        FROM papers p LEFT JOIN users u ON u.id = p.created_by
        ORDER BY p.id DESC LIMIT 5
        """
    ).fetchall()
    return render_template(
        "dashboard.html",
        stats=stats,
        new_counts=new_counts,
        recent_papers=recent_papers,
        student_class=student_class,
    )


@app.get("/classes")
@login_required
@teacher_required
def classes():
    rows = get_db().execute(
        """
        SELECT c.id, c.name, c.created_at,
               COUNT(u.id) AS student_count
        FROM classes c
        LEFT JOIN users u ON u.class_id = c.id AND u.role = 'student'
        WHERE c.teacher_id = ?
        GROUP BY c.id
        ORDER BY c.id DESC
        """,
        (session["user"]["id"],),
    ).fetchall()
    return render_template("classes.html", classes=rows)


@app.route("/classes/new", methods=("GET", "POST"))
@login_required
@teacher_required
def class_new():
    if request.method == "POST":
        name = request.form.get("name", "").strip()
        if not name:
            flash("班级名称不能为空。", "error")
        else:
            get_db().execute(
                "INSERT INTO classes (name, teacher_id) VALUES (?, ?)",
                (name, session["user"]["id"]),
            )
            get_db().commit()
            flash("班级已创建。", "success")
            return redirect(url_for("classes"))
    return render_template("class_form.html", class_item=None)


@app.route("/classes/<int:class_id>/edit", methods=("GET", "POST"))
@login_required
@teacher_required
def class_edit(class_id):
    db = get_db()
    class_item = db.execute(
        "SELECT * FROM classes WHERE id = ? AND teacher_id = ?",
        (class_id, session["user"]["id"]),
    ).fetchone()
    if class_item is None:
        flash("班级不存在。", "error")
        return redirect(url_for("classes"))
    if request.method == "POST":
        name = request.form.get("name", "").strip()
        if not name:
            flash("班级名称不能为空。", "error")
        else:
            db.execute(
                "UPDATE classes SET name = ? WHERE id = ? AND teacher_id = ?",
                (name, class_id, session["user"]["id"]),
            )
            db.commit()
            flash("班级名称已更新。", "success")
            return redirect(url_for("classes"))
    return render_template("class_form.html", class_item=class_item)


@app.post("/classes/<int:class_id>/delete")
@login_required
@teacher_required
def class_delete(class_id):
    db = get_db()
    db.execute(
        "UPDATE users SET class_id = NULL WHERE class_id = ? AND role = 'student'",
        (class_id,),
    )
    db.execute(
        "DELETE FROM classes WHERE id = ? AND teacher_id = ?",
        (class_id, session["user"]["id"]),
    )
    db.commit()
    flash("班级已删除，原班级学生已变为未分班。", "success")
    return redirect(url_for("classes"))


@app.get("/classes/<int:class_id>")
@login_required
@teacher_required
def class_detail(class_id):
    db = get_db()
    class_item = db.execute(
        "SELECT * FROM classes WHERE id = ? AND teacher_id = ?",
        (class_id, session["user"]["id"]),
    ).fetchone()
    if class_item is None:
        flash("班级不存在。", "error")
        return redirect(url_for("classes"))
    students = db.execute(
        """
        SELECT id, username, display_name
        FROM users
        WHERE role = 'student' AND class_id = ?
        ORDER BY id ASC
        """,
        (class_id,),
    ).fetchall()
    unassigned_students = db.execute(
        """
        SELECT id, username, display_name
        FROM users
        WHERE role = 'student' AND (class_id IS NULL OR class_id != ?)
        ORDER BY id ASC
        """,
        (class_id,),
    ).fetchall()
    return render_template(
        "class_detail.html",
        class_item=class_item,
        students=students,
        unassigned_students=unassigned_students,
    )


@app.post("/classes/<int:class_id>/students")
@login_required
@teacher_required
def class_add_student(class_id):
    db = get_db()
    class_item = db.execute(
        "SELECT * FROM classes WHERE id = ? AND teacher_id = ?",
        (class_id, session["user"]["id"]),
    ).fetchone()
    if class_item is None:
        flash("班级不存在。", "error")
        return redirect(url_for("classes"))

    lookup = request.form.get("student_lookup", "").strip()
    if not lookup:
        flash("请输入学生账号或姓名。", "error")
        return redirect(url_for("class_detail", class_id=class_id))

    students = db.execute(
        """
        SELECT * FROM users
        WHERE role = 'student'
          AND (username = ? OR display_name = ? OR (display_name || '（' || username || '）') = ?)
        ORDER BY id ASC
        """,
        (lookup, lookup, lookup),
    ).fetchall()
    if not students:
        flash("没有找到对应学生，请检查账号或姓名。", "error")
        return redirect(url_for("class_detail", class_id=class_id))
    if len(students) > 1:
        flash("找到多个同名或近似学生，请使用完整账号添加。", "error")
        return redirect(url_for("class_detail", class_id=class_id))

    student = students[0]
    db.execute(
        "UPDATE users SET class_id = ? WHERE id = ? AND role = 'student'",
        (class_id, student["id"]),
    )
    db.commit()
    flash(f"{student['display_name'] or student['username']} 已加入班级。", "success")
    return redirect(url_for("class_detail", class_id=class_id))


@app.post("/classes/<int:class_id>/students/<int:user_id>/remove")
@login_required
@teacher_required
def class_remove_student(class_id, user_id):
    db = get_db()
    db.execute(
        "UPDATE users SET class_id = NULL WHERE id = ? AND role = 'student' AND class_id = ?",
        (user_id, class_id),
    )
    db.commit()
    flash("学生已移出班级。", "success")
    return redirect(url_for("class_detail", class_id=class_id))


@app.route("/classes/<int:class_id>/import", methods=("GET", "POST"))
@login_required
@teacher_required
def class_student_import(class_id):
    db = get_db()
    class_item = db.execute(
        "SELECT * FROM classes WHERE id = ? AND teacher_id = ?",
        (class_id, session["user"]["id"]),
    ).fetchone()
    if class_item is None:
        flash("班级不存在。", "error")
        return redirect(url_for("classes"))

    result = None
    if request.method == "POST":
        uploaded_file = request.files.get("file")
        if not uploaded_file or not uploaded_file.filename:
            flash("请选择 CSV 或 Excel 文件。", "error")
            return render_template("class_student_import.html", class_item=class_item, result=result)

        try:
            filename = uploaded_file.filename.lower()
            if filename.endswith((".xlsx", ".xlsm")):
                raw_rows = _rows_from_excel(uploaded_file.read())
            else:
                text = _decode_import_upload(uploaded_file)
                raw_rows = [
                    row
                    for row in csv.DictReader(io.StringIO(text, newline=""))
                    if any((row.get(key) or "").strip() for key in row)
                ]
        except Exception:
            flash("文件读取失败，请确认是 CSV 或 Excel 文件。", "error")
            return render_template("class_student_import.html", class_item=class_item, result=result)

        added = 0
        updated = 0
        skipped = 0
        row_errors = []
        for index, raw_row in enumerate(raw_rows, start=2):
            name = ""
            username = ""
            for key, value in raw_row.items():
                normalized_key = _import_text(key)
                if normalized_key in ("姓名", "学生姓名", "name", "display_name"):
                    name = _import_text(value)
                elif normalized_key in ("账号", "学号", "用户名", "username", "account"):
                    username = _import_text(value)
            if not username:
                row_errors.append(f"第 {index} 行：账号不能为空。")
                continue
            existing = db.execute(
                "SELECT * FROM users WHERE username = ?",
                (username,),
            ).fetchone()
            if existing:
                if existing["role"] != "student":
                    row_errors.append(f"第 {index} 行：账号 {username} 不是学生账号。")
                    continue
                db.execute(
                    """
                    UPDATE users
                    SET class_id = ?, display_name = CASE WHEN ? != '' THEN ? ELSE display_name END
                    WHERE id = ?
                    """,
                    (class_id, name, name, existing["id"]),
                )
                updated += 1
            else:
                db.execute(
                    """
                    INSERT INTO users (username, display_name, password_hash, role, class_id)
                    VALUES (?, ?, ?, 'student', ?)
                    """,
                    (
                        username,
                        name or username,
                        generate_password_hash(f"czu{username}"),
                        class_id,
                    ),
                )
                added += 1
        db.commit()
        result = {"added": added, "updated": updated, "skipped": skipped, "errors": row_errors}
        if added or updated:
            flash(f"导入完成：新增 {added} 人，更新 {updated} 人。", "success")
        else:
            flash("没有成功导入学生。", "error")

    return render_template("class_student_import.html", class_item=class_item, result=result)


@app.get("/questions")
@login_required
def questions():
    keyword = request.args.get("keyword", "").strip()
    subject = request.args.get("subject", "").strip()
    difficulty = request.args.get("difficulty", "").strip()
    question_type = request.args.get("question_type", "").strip()
    page = request.args.get("page", 1, type=int)
    per_page = 20
    where_clause = "1 = 1"
    params = []
    if keyword:
        where_clause += " AND (content LIKE ? OR chapter LIKE ? OR subject LIKE ?)"
        params.extend([f"%{keyword}%", f"%{keyword}%", f"%{keyword}%"])
    if subject:
        where_clause += " AND subject = ?"
        params.append(subject)
    if difficulty in DIFFICULTY_LABELS:
        where_clause += " AND difficulty = ?"
        params.append(difficulty)
    if question_type in QUESTION_TYPE_LABELS:
        where_clause += " AND question_type = ?"
        params.append(question_type)
    total = get_db().execute(f"SELECT COUNT(*) AS count FROM questions WHERE {where_clause}", params).fetchone()["count"]
    total_pages = max(1, (total + per_page - 1) // per_page)
    page = min(max(page, 1), total_pages)
    offset = (page - 1) * per_page
    rows = get_db().execute(
        f"SELECT * FROM questions WHERE {where_clause} ORDER BY id DESC LIMIT ? OFFSET ?",
        [*params, per_page, offset],
    ).fetchall()
    return render_template(
        "questions.html",
        questions=rows,
        filters={
            "keyword": keyword,
            "subject": subject,
            "difficulty": difficulty,
            "question_type": question_type,
        },
        pagination={
            "page": page,
            "total_pages": total_pages,
            "total": total,
            "per_page": per_page,
        },
    )


@app.route("/questions/new", methods=("GET", "POST"))
@login_required
@teacher_required
def question_new():
    if request.method == "POST":
        form = question_form_data()
        error = validate_question_form(form)
        if error:
            flash(error, "error")
        else:
            get_db().execute(
                """
                INSERT INTO questions
                (content, question_type, difficulty, subject, chapter, options, answer, analysis)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                """,
                tuple(form.values()),
            )
            get_db().commit()
            flash("题目已添加。", "success")
            return redirect(url_for("questions"))
    else:
        form = empty_question_form()
    return render_template("question_form.html", form=form, editing=False)


@app.route("/questions/<int:question_id>/edit", methods=("GET", "POST"))
@login_required
@teacher_required
def question_edit(question_id):
    db = get_db()
    question = db.execute("SELECT * FROM questions WHERE id = ?", (question_id,)).fetchone()
    if question is None:
        flash("题目不存在。", "error")
        return redirect(url_for("questions"))
    if request.method == "POST":
        form = question_form_data()
        error = validate_question_form(form)
        if error:
            flash(error, "error")
        else:
            db.execute(
                """
                UPDATE questions
                SET content = ?, question_type = ?, difficulty = ?, subject = ?, chapter = ?, options = ?, answer = ?, analysis = ?
                WHERE id = ?
                """,
                (*form.values(), question_id),
            )
            db.commit()
            flash("题目已更新。", "success")
            return redirect(url_for("questions"))
    else:
        form = {key: question[key] for key in empty_question_form()}
    return render_template("question_form.html", form=form, editing=True, question_id=question_id)


@app.post("/questions/<int:question_id>/delete")
@login_required
@teacher_required
def question_delete(question_id):
    db = get_db()
    db.execute("DELETE FROM questions WHERE id = ?", (question_id,))
    db.commit()
    flash("题目已删除。", "success")
    return redirect(url_for("questions"))


QUESTION_IMPORT_FIELDS = ("subject", "chapter", "question_type", "difficulty", "content", "options", "answer", "analysis")
QUESTION_IMPORT_HEADERS = ("科目", "章节", "题型", "难度", "题目内容", "选项", "答案", "解析")
QUESTION_IMPORT_ALIASES = {
    "科目": "subject", "章节": "chapter", "题型": "question_type", "难度": "difficulty",
    "题目内容": "content", "选项": "options", "答案": "answer", "解析": "analysis",
    "subject": "subject", "chapter": "chapter", "question_type": "question_type", "difficulty": "difficulty",
    "content": "content", "options": "options", "answer": "answer", "analysis": "analysis",
}


def _import_text(value):
    if value is None:
        return ""
    return str(value).strip()


def _import_question_type(value):
    value = _import_text(value)
    for key, label in QUESTION_TYPE_LABELS.items():
        if value in (key, label):
            return key
    return value


def _import_difficulty(value):
    value = _import_text(value)
    for key, label in DIFFICULTY_LABELS.items():
        if value in (key, label):
            return key
    return value


def _import_options(value):
    value = _import_text(value).replace("\r\n", "\n").replace("\r", "\n")
    if any(separator in value for separator in ("|", "｜", "；", ";")):
        parts = [part.strip() for part in re.split(r"[|｜；;\n]", value) if part.strip()]
        return "\n".join(parts)
    return value


def _normalize_import_row(row):
    normalized = {}
    for raw_key, value in row.items():
        key = QUESTION_IMPORT_ALIASES.get(_import_text(raw_key))
        if key not in QUESTION_IMPORT_FIELDS:
            continue
        normalized[key] = _import_text(value)
    for field in QUESTION_IMPORT_FIELDS:
        normalized.setdefault(field, "")
    normalized["question_type"] = _import_question_type(normalized["question_type"])
    normalized["difficulty"] = _import_difficulty(normalized["difficulty"])
    normalized["options"] = _import_options(normalized["options"])
    if not normalized["subject"]:
        normalized["subject"] = DEFAULT_SUBJECT
    return normalized


def _decode_import_upload(uploaded_file):
    raw = uploaded_file.read()
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _rows_from_excel(raw):
    workbook = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    sheet = workbook.active
    rows_iter = sheet.iter_rows(values_only=True)
    try:
        header_values = next(rows_iter)
    except StopIteration:
        return []
    headers = [_import_text(value) or f"__col_{index}" for index, value in enumerate(header_values, start=1)]
    rows = []
    for values in rows_iter:
        if not any(value is not None and _import_text(value) for value in values):
            continue
        rows.append(dict(zip(headers, values)))
    return rows


@app.get("/questions/import/template")
@login_required
@teacher_required
def question_import_template():
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(QUESTION_IMPORT_HEADERS)
    writer.writerow([
        "计算机基础",
        "Python基础",
        "单选题",
        "简单",
        "Python 中用于定义函数的关键字是什么？",
        "A. def\nB. class\nC. import\nD. return",
        "A",
        "def 用于定义函数。",
    ])
    data = output.getvalue().encode("utf-8-sig")
    return send_file(
        io.BytesIO(data),
        mimetype="text/csv",
        as_attachment=True,
        download_name="题目批量导入模板.csv",
    )


@app.get("/questions/import/template/xlsx")
@login_required
@teacher_required
def question_import_excel_template():
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "题目导入模板"
    sheet.append(QUESTION_IMPORT_HEADERS)
    sheet.append([
        "计算机基础",
        "Python基础",
        "单选题",
        "简单",
        "Python 中用于定义函数的关键字是什么？",
        "A. def\nB. class\nC. import\nD. return",
        "A",
        "def 用于定义函数。",
    ])
    for cell in sheet[1]:
        cell.font = openpyxl.styles.Font(bold=True)
    for column_cells in sheet.columns:
        width = max((len(_import_text(cell.value)) for cell in column_cells if cell.value is not None), default=10)
        sheet.column_dimensions[column_cells[0].column_letter].width = min(max(width + 2, 12), 42)

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name="题目批量导入模板.xlsx",
    )


@app.route("/questions/import", methods=("GET", "POST"))
@login_required
@teacher_required
def question_import():
    result = None
    if request.method == "POST":
        uploaded_file = request.files.get("file")
        if not uploaded_file or not uploaded_file.filename:
            flash("请选择要导入的 CSV 或 Excel 文件。", "error")
            return render_template("question_import.html", result=result)

        try:
            filename = uploaded_file.filename.lower()
            if filename.endswith((".xlsx", ".xlsm")):
                rows = _rows_from_excel(uploaded_file.read())
            else:
                text = _decode_import_upload(uploaded_file)
                rows = [
                    row
                    for row in csv.DictReader(io.StringIO(text, newline=""))
                    if any((row.get(key) or "").strip() for key in row)
                ]
        except Exception:
            flash("文件读取失败，请确认是 UTF-8/GBK 编码的 CSV，或有效的 .xlsx 文件。", "error")
            return render_template("question_import.html", result=result)

        if not rows:
            flash("文件中没有可导入的数据行。", "error")
            return render_template("question_import.html", result=result)

        db = get_db()
        inserted = 0
        skipped = 0
        row_errors = []

        for index, row in enumerate(rows, start=2):
            normalized = _normalize_import_row(row)
            if normalized["subject"] not in SUBJECTS:
                row_errors.append(f"第 {index} 行：科目“{normalized['subject']}”不在支持的科目列表中。")
                continue
            error = validate_question_form(normalized)
            if error:
                row_errors.append(f"第 {index} 行：{error}")
                continue
            exists = db.execute(
                "SELECT 1 FROM questions WHERE subject = ? AND content = ?",
                (normalized["subject"], normalized["content"]),
            ).fetchone()
            if exists:
                skipped += 1
                continue
            db.execute(
                """
                INSERT INTO questions
                (content, question_type, difficulty, subject, chapter, options, answer, analysis, source)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, 'import')
                """,
                (
                    normalized["content"],
                    normalized["question_type"],
                    normalized["difficulty"],
                    normalized["subject"],
                    normalized["chapter"],
                    normalized["options"],
                    normalized["answer"],
                    normalized["analysis"],
                ),
            )
            inserted += 1

        if inserted:
            db.commit()

        result = {
            "total": len(rows),
            "inserted": inserted,
            "skipped": skipped,
            "errors": row_errors,
        }

        if inserted and not row_errors:
            flash(f"成功导入 {inserted} 道题目。", "success")
            return redirect(url_for("questions"))
        if inserted:
            flash(f"部分导入成功：新增 {inserted} 道，跳过重复 {skipped} 道。", "warning")
        else:
            flash(f"没有成功导入题目，共发现 {len(row_errors)} 个问题。", "error")

    return render_template("question_import.html", result=result)


@app.route("/papers/new", methods=("GET", "POST"))
@login_required
@teacher_required
def paper_new():
    selected_subject = request.values.get("subject", DEFAULT_SUBJECT).strip() or DEFAULT_SUBJECT
    chapters = get_chapters(selected_subject)
    subject_chapters = get_subject_chapters()
    availability = get_question_availability()
    teacher_classes = get_db().execute(
        """
        SELECT id, name FROM classes
        WHERE teacher_id = ?
        ORDER BY id DESC
        """,
        (session["user"]["id"],),
    ).fetchall()
    selected_class_ids = request.form.getlist("class_ids") if request.method == "POST" else []
    preview_draft = None
    replacement_choices = []
    action = request.form.get("action", "preview") if request.method == "POST" else ""
    form_data = {"subject": selected_subject}
    if request.method == "POST":
        form_data = dict(request.form)
        title = request.form.get("title", "").strip()
        if not title:
            flash("试卷名称不能为空。", "error")
            return render_template(
                "paper_form.html",
                form=form_data,
                chapters=chapters,
                subject_chapters=subject_chapters,
                availability=availability,
                selected_subject=selected_subject,
                preview_draft=preview_draft,
                replacement_choices=replacement_choices,
                teacher_classes=teacher_classes,
                selected_class_ids=selected_class_ids,
            )
        result = build_paper_from_form(request.form)
        if isinstance(result, str):
            flash(result, "error")
            return render_template(
                "paper_form.html",
                form=form_data,
                chapters=chapters,
                subject_chapters=subject_chapters,
                availability=availability,
                selected_subject=selected_subject,
                preview_draft=preview_draft,
                replacement_choices=replacement_choices,
                teacher_classes=teacher_classes,
                selected_class_ids=selected_class_ids,
            )
        preview_draft = result
        replacement_choices = build_replacement_choices(result["questions"], selected_subject, request.form.get("chapter", "").strip())
        if action == "generate":
            applied = apply_preview_replacements(result["questions"], request.form, selected_subject, request.form.get("chapter", "").strip())
            if isinstance(applied, str):
                flash(applied, "error")
                preview_draft = result
            else:
                class_ids = [int(value) for value in request.form.getlist("class_ids") if value.isdigit()]
                if not class_ids:
                    flash("请至少选择一个发布班级。", "error")
                    preview_draft = result
                else:
                    result["questions"] = applied
                    paper_id = save_paper(session["user"]["id"], title, result, class_ids=class_ids)
                    flash("试卷已生成。", "success")
                    return redirect(url_for("paper_detail", paper_id=paper_id))
        else:
            flash("已生成预览，可在下方更换题目后再生成。", "success")
    return render_template(
        "paper_form.html",
        form=form_data,
        chapters=chapters,
        subject_chapters=subject_chapters,
        availability=availability,
        selected_subject=selected_subject,
        preview_draft=preview_draft,
        replacement_choices=replacement_choices,
        teacher_classes=teacher_classes,
        selected_class_ids=selected_class_ids,
    )


@app.get("/papers/<int:paper_id>")
@login_required
def paper_detail(paper_id):
    if not current_user_is_teacher():
        flash("学生账号不能预览试卷题目，请从考试入口进入。", "error")
        return redirect(url_for("take_paper", paper_id=paper_id))

    db = get_db()
    paper = db.execute(
        """
        SELECT p.*, COALESCE(NULLIF(u.display_name, ''), u.username) AS creator FROM papers p
        LEFT JOIN users u ON u.id = p.created_by WHERE p.id = ?
        """,
        (paper_id,),
    ).fetchone()
    if paper is None:
        flash("试卷不存在。", "error")
        return redirect(url_for("dashboard"))
    items = get_paper_items(paper_id)
    paper_subject = get_paper_subject(paper_id)
    distribution = build_paper_distribution(items)
    return render_template(
        "paper_detail.html",
        paper=paper,
        items=items,
        paper_subject=paper_subject,
        distribution=distribution,
        preview=current_user_is_teacher(),
    )


@app.get("/papers")
@login_required
def papers():
    db = get_db()
    selected_class_id = request.args.get("class_id", "").strip()
    class_options = []
    if current_user_is_teacher():
        class_options = db.execute(
            """
            SELECT id, name FROM classes
            WHERE teacher_id = ?
            ORDER BY id DESC
            """,
            (session["user"]["id"],),
        ).fetchall()
    if current_user_is_teacher():
        rows = db.execute(
            """
            SELECT p.*, COALESCE(NULLIF(u.display_name, ''), u.username) AS creator
            FROM papers p
            LEFT JOIN users u ON u.id = p.created_by
            ORDER BY p.id DESC
            """
        ).fetchall()
    else:
        student_class = db.execute(
            "SELECT class_id FROM users WHERE id = ?",
            (session["user"]["id"],),
        ).fetchone()
        student_class_id = student_class["class_id"] if student_class else None
        rows = db.execute(
            """
            SELECT p.*, COALESCE(NULLIF(u.display_name, ''), u.username) AS creator,
                   CASE WHEN a.id IS NULL THEN 'not_started' ELSE 'completed' END AS status,
                   a.id AS attempt_id,
                   COALESCE(ac.attempt_count, 0) AS attempt_count
            FROM papers p
            LEFT JOIN users u ON u.id = p.created_by
            LEFT JOIN (
                SELECT paper_id, COUNT(*) AS attempt_count
                FROM exam_attempts
                WHERE user_id = ?
                GROUP BY paper_id
            ) ac ON ac.paper_id = p.id
            LEFT JOIN (
                SELECT a1.*
                FROM exam_attempts a1
                JOIN (
                    SELECT paper_id, MAX(id) AS id
                    FROM exam_attempts
                    WHERE user_id = ?
                    GROUP BY paper_id
                ) latest ON latest.id = a1.id
            ) a ON a.paper_id = p.id
            WHERE (
                p.id NOT IN (SELECT paper_id FROM paper_classes)
                OR p.id IN (SELECT paper_id FROM paper_classes WHERE class_id = ?)
            )
            ORDER BY p.id DESC
            """,
            (session["user"]["id"], session["user"]["id"], student_class_id),
        ).fetchall()
    total_papers = len(rows)
    papers = []
    for index, row in enumerate(rows):
        paper = {
            **dict(row),
            "subject": paper_subjects(get_paper_items(row["id"])),
            "display_id": total_papers - index,
        }
        if not current_user_is_teacher():
            paper["remaining_attempts"] = max(0, paper.get("max_attempts", 1) - paper.get("attempt_count", 0))
        else:
            if selected_class_id.isdigit():
                class_id = int(selected_class_id)
                total_students = db.execute(
                    "SELECT COUNT(*) AS count FROM users WHERE role = 'student' AND class_id = ?",
                    (class_id,),
                ).fetchone()["count"]
                completed_students = db.execute(
                    """
                    SELECT COUNT(DISTINCT a.user_id) AS count
                    FROM exam_attempts a
                    JOIN users u ON u.id = a.user_id
                    WHERE a.paper_id = ? AND u.class_id = ?
                    """,
                    (paper["id"], class_id),
                ).fetchone()["count"]
                paper["completion_text"] = f"{completed_students}/{total_students}"
            else:
                paper["completion_text"] = "—"
        papers.append(paper)
    return render_template(
        "papers.html",
        papers=papers,
        class_options=class_options,
        selected_class_id=selected_class_id,
    )


@app.get("/papers/<int:paper_id>/export/<file_type>")
@login_required
@teacher_required
def export_paper(paper_id, file_type):
    paper = get_paper(paper_id)
    if paper is None:
        flash("试卷不存在。", "error")
        return redirect(url_for("papers"))

    items = get_paper_items(paper_id)
    if not items:
        flash("这份试卷没有题目，暂时不能导出。", "error")
        return redirect(url_for("paper_detail", paper_id=paper_id))

    if file_type == "docx":
        export_path = build_paper_docx(paper, items)
    elif file_type == "pdf":
        export_path = build_paper_pdf(paper, items)
    else:
        flash("不支持的导出格式。", "error")
        return redirect(url_for("paper_detail", paper_id=paper_id))

    flash(f"试卷已导出到 exports 文件夹：{export_path.name}", "success")
    return send_file(export_path, as_attachment=True, download_name=export_path.name)


@app.get("/results")
@login_required
def results():
    db = get_db()
    filters = {
        "student_keyword": request.args.get("student_keyword", "").strip(),
        "paper_keyword": request.args.get("paper_keyword", "").strip(),
        "min_score": request.args.get("min_score", "").strip(),
        "max_score": request.args.get("max_score", "").strip(),
    }
    min_score, min_error = parse_optional_score(filters["min_score"], "最低分")
    max_score, max_error = parse_optional_score(filters["max_score"], "最高分")
    if min_error:
        flash(min_error, "error")
    if max_error:
        flash(max_error, "error")
    if min_score is not None and max_score is not None and min_score > max_score:
        flash("最低分不能大于最高分。", "error")
        min_score = None
        max_score = None

    sql = """
        SELECT a.*, p.title, u.username AS account,
               COALESCE(NULLIF(u.display_name, ''), u.username) AS username,
               ROUND(a.score * 100.0 / NULLIF(a.total_score, 0), 1) AS score_rate
        FROM exam_attempts a
        JOIN papers p ON p.id = a.paper_id
        JOIN users u ON u.id = a.user_id
        WHERE 1 = 1
    """
    params = []
    if current_user_is_teacher():
        if filters["student_keyword"]:
            sql += " AND (u.username LIKE ? OR u.display_name LIKE ?)"
            params.append(f"%{filters['student_keyword']}%")
            params.append(f"%{filters['student_keyword']}%")
    else:
        sql += " AND a.user_id = ?"
        params.append(session["user"]["id"])

    if filters["paper_keyword"]:
        sql += " AND p.title LIKE ?"
        params.append(f"%{filters['paper_keyword']}%")
    if min_score is not None:
        sql += " AND a.score >= ?"
        params.append(min_score)
    if max_score is not None:
        sql += " AND a.score <= ?"
        params.append(max_score)

    sql += " ORDER BY a.submitted_at DESC, a.id DESC"
    attempts = db.execute(sql, params).fetchall()
    stats = build_result_stats(attempts)
    return render_template(
        "results.html",
        attempts=attempts,
        filters=filters,
        stats=stats,
    )


@app.get("/analytics")
@login_required
def analytics():
    db = get_db()
    user_filter = "" if current_user_is_teacher() else "AND a.user_id = ?"
    params = [] if current_user_is_teacher() else [session["user"]["id"]]
    class_options = []
    if current_user_is_teacher():
        class_options = db.execute(
            """
            SELECT id, name FROM classes
            WHERE teacher_id = ?
            ORDER BY id DESC
            """,
            (session["user"]["id"],),
        ).fetchall()
    subject_options = db.execute(
        """
        SELECT DISTINCT COALESCE(NULLIF(subject, ''), ?) AS subject
        FROM questions
        ORDER BY subject
        """,
        (DEFAULT_SUBJECT,),
    ).fetchall()
    paper_options = db.execute(
        """
        SELECT p.id, p.title
        FROM papers p
        ORDER BY p.id DESC
        """
    ).fetchall()
    selected_subject = request.args.get("subject", "").strip()
    if not current_user_is_teacher() and not selected_subject:
        selected_subject = first_student_subject(session["user"]["id"])
    selected_paper_id = request.args.get("paper_id", "").strip()
    selected_class_id = request.args.get("class_id", "").strip()
    selected_paper_title = ""
    if selected_paper_id.isdigit():
        paper_row = db.execute("SELECT title FROM papers WHERE id = ?", (int(selected_paper_id),)).fetchone()
        if paper_row:
            selected_paper_title = paper_row["title"]
            user_filter += " AND a.paper_id = ?"
            params.append(int(selected_paper_id))
    if selected_subject:
        user_filter += """
            AND EXISTS (
                SELECT 1
                FROM paper_questions fpq
                JOIN questions fq ON fq.id = fpq.question_id
                WHERE fpq.paper_id = a.paper_id AND fq.subject = ?
            )
        """
        params.append(selected_subject)
    if current_user_is_teacher() and selected_class_id.isdigit():
        class_id = int(selected_class_id)
        user_filter += " AND a.user_id IN (SELECT id FROM users WHERE class_id = ?)"
        params.append(class_id)
    class_analysis = None
    class_students = []
    if current_user_is_teacher() and selected_class_id.isdigit():
        class_id = int(selected_class_id)
        class_analysis = db.execute(
            """
            SELECT c.name,
                   COUNT(DISTINCT a.id) AS attempts,
                   COALESCE(ROUND(AVG(a.score * 100.0 / NULLIF(a.total_score, 0)), 1), 0) AS avg_rate,
                   COALESCE(MAX(a.score), 0) AS best_score,
                   COUNT(DISTINCT u.id) AS student_count
            FROM classes c
            LEFT JOIN users u ON u.class_id = c.id AND u.role = 'student'
            LEFT JOIN exam_attempts a ON a.user_id = u.id
            WHERE c.id = ? AND c.teacher_id = ?
            GROUP BY c.id
            """,
            (class_id, session["user"]["id"]),
        ).fetchone()
        class_students = db.execute(
            """
            SELECT u.id, u.username, u.display_name,
                   COUNT(a.id) AS attempts,
                   COALESCE(ROUND(AVG(a.score * 100.0 / NULLIF(a.total_score, 0)), 1), 0) AS avg_rate,
                   COALESCE(MAX(a.score), 0) AS best_score,
                   MAX(a.submitted_at) AS last_submitted
            FROM users u
            LEFT JOIN exam_attempts a ON a.user_id = u.id
            WHERE u.role = 'student' AND u.class_id = ?
            GROUP BY u.id
            ORDER BY u.username ASC
            """,
            (class_id,),
        ).fetchall()
    overview = db.execute(
        f"""
        SELECT COUNT(DISTINCT a.id) AS attempts,
               COALESCE(ROUND(AVG(a.score * 100.0 / NULLIF(a.total_score, 0)), 1), 0) AS avg_rate,
               COALESCE(MAX(a.score), 0) AS best_score
        FROM exam_attempts a
        WHERE 1 = 1 {user_filter}
        """,
        params,
    ).fetchone()
    recent = db.execute(
        f"""
        SELECT a.score, a.total_score, a.submitted_at, p.title, u.username AS account,
               COALESCE(NULLIF(u.display_name, ''), u.username) AS username
        FROM exam_attempts a
        JOIN papers p ON p.id = a.paper_id
        JOIN users u ON u.id = a.user_id
        WHERE 1 = 1 {user_filter}
        ORDER BY a.submitted_at DESC, a.id DESC
        LIMIT 8
        """,
        params,
    ).fetchall()
    score_rows = db.execute(
        f"""
        SELECT a.score, a.total_score
        FROM exam_attempts a
        WHERE 1 = 1 {user_filter}
        """,
        params,
    ).fetchall()
    chart_data = {
        "chapter": analytics_group("q.chapter", user_filter, params),
        "question_type": analytics_group("q.question_type", user_filter, params, QUESTION_TYPE_LABELS),
        "difficulty": analytics_group("q.difficulty", user_filter, params, DIFFICULTY_LABELS),
    }
    score_buckets = build_score_buckets(score_rows)
    score_curve = build_score_curve(score_buckets)
    ability_radar_data = []
    ability_trend_data = []
    if not current_user_is_teacher():
        ability_radar_data = build_ability_radar_data(session["user"]["id"])
        ability_trend_data = build_ability_trend(session["user"]["id"], selected_subject)
    return render_template(
        "analytics.html",
        overview=overview,
        recent=recent,
        chart_data=chart_data,
        score_buckets=score_buckets,
        score_curve=score_curve,
        subject_options=subject_options,
        paper_options=paper_options,
        class_options=class_options,
        selected_subject=selected_subject,
        selected_paper_id=selected_paper_id,
        selected_class_id=selected_class_id,
        selected_paper_title=selected_paper_title,
        class_analysis=class_analysis,
        class_students=class_students,
        ability_radar_data=ability_radar_data,
        ability_trend_data=ability_trend_data,
    )


@app.get("/analytics/student/<int:user_id>")
@login_required
@teacher_required
def student_analytics(user_id):
    db = get_db()
    student = db.execute(
        """
        SELECT u.id, u.username, u.display_name, u.class_id,
               c.name AS class_name, c.teacher_id
        FROM users u
        LEFT JOIN classes c ON c.id = u.class_id
        WHERE u.id = ? AND u.role = 'student'
        """,
        (user_id,),
    ).fetchone()
    if student is None:
        flash("学生不存在。", "error")
        return redirect(url_for("analytics"))
    if student["class_id"] and student["teacher_id"] != session["user"]["id"]:
        flash("你无权查看这个班级的学生。", "error")
        return redirect(url_for("analytics"))

    overview = db.execute(
        """
        SELECT COUNT(*) AS attempts,
               COALESCE(ROUND(AVG(score * 100.0 / NULLIF(total_score, 0)), 1), 0) AS avg_rate,
               COALESCE(MAX(score), 0) AS best_score
        FROM exam_attempts
        WHERE user_id = ?
        """,
        (user_id,),
    ).fetchone()
    recent = db.execute(
        """
        SELECT a.id, a.score, a.total_score, a.submitted_at, p.title,
               ROUND(a.score * 100.0 / NULLIF(a.total_score, 0), 1) AS score_rate
        FROM exam_attempts a
        JOIN papers p ON p.id = a.paper_id
        WHERE a.user_id = ?
        ORDER BY a.submitted_at DESC, a.id DESC
        LIMIT 8
        """,
        (user_id,),
    ).fetchall()
    chart_data = {
        "chapter": analytics_group("q.chapter", "AND a.user_id = ?", [user_id]),
        "question_type": analytics_group("q.question_type", "AND a.user_id = ?", [user_id], QUESTION_TYPE_LABELS),
        "difficulty": analytics_group("q.difficulty", "AND a.user_id = ?", [user_id], DIFFICULTY_LABELS),
    }
    ability_radar_data = build_ability_radar_data(user_id)
    ability_trend_data = build_ability_trend(user_id)
    return render_template(
        "student_analytics.html",
        student=student,
        overview=overview,
        recent=recent,
        chart_data=chart_data,
        ability_radar_data=ability_radar_data,
        ability_trend_data=ability_trend_data,
    )


@app.get("/wrong-questions")
@login_required
def wrong_questions():
    keyword = request.args.get("keyword", "").strip()
    subject = request.args.get("subject", "").strip()
    difficulty = request.args.get("difficulty", "").strip()
    question_type = request.args.get("question_type", "").strip()

    sql = """
        SELECT ea.answer AS user_answer, ea.score, a.submitted_at, p.title,
               q.content, q.answer AS correct_answer, q.analysis, q.chapter,
               q.question_type, q.difficulty, q.subject
        FROM exam_answers ea
        JOIN exam_attempts a ON a.id = ea.attempt_id
        JOIN papers p ON p.id = a.paper_id
        JOIN paper_questions pq ON pq.id = ea.paper_question_id
        JOIN questions q ON q.id = pq.question_id
        WHERE a.user_id = ? AND ea.is_correct = 0
    """
    params = [session["user"]["id"]]
    if keyword:
        sql += " AND (q.content LIKE ? OR q.chapter LIKE ? OR p.title LIKE ?)"
        params.extend([f"%{keyword}%", f"%{keyword}%", f"%{keyword}%"])
    if subject:
        sql += " AND q.subject = ?"
        params.append(subject)
    if difficulty in DIFFICULTY_LABELS:
        sql += " AND q.difficulty = ?"
        params.append(difficulty)
    if question_type in QUESTION_TYPE_LABELS:
        sql += " AND q.question_type = ?"
        params.append(question_type)
    sql += " ORDER BY a.submitted_at DESC, ea.id DESC"

    rows = get_db().execute(sql, params).fetchall()
    return render_template(
        "wrong_questions.html",
        wrong_questions=rows,
        filters={
            "keyword": keyword,
            "subject": subject,
            "difficulty": difficulty,
            "question_type": question_type,
        },
    )


@app.route("/papers/<int:paper_id>/take", methods=("GET", "POST"))
@login_required
def take_paper(paper_id):
    if current_user_is_teacher():
        return redirect(url_for("paper_detail", paper_id=paper_id))

    db = get_db()
    paper = db.execute("SELECT * FROM papers WHERE id = ?", (paper_id,)).fetchone()
    if paper is None:
        flash("试卷不存在。", "error")
        return redirect(url_for("papers"))
    student_class = db.execute(
        "SELECT class_id FROM users WHERE id = ?",
        (session["user"]["id"],),
    ).fetchone()
    student_class_id = student_class["class_id"] if student_class else None
    visible = db.execute(
        """
        SELECT 1 FROM papers p
        WHERE p.id = ?
          AND (
            p.id NOT IN (SELECT paper_id FROM paper_classes)
            OR p.id IN (SELECT paper_id FROM paper_classes WHERE class_id = ?)
          )
        """,
        (paper_id, student_class_id),
    ).fetchone()
    if visible is None:
        flash("这份试卷未发布到你的班级，暂时不能参加考试。", "error")
        return redirect(url_for("papers"))
    items = get_paper_items(paper_id)
    if not items:
        flash("这份试卷没有题目，暂时不能开始考试。", "error")
        return redirect(url_for("papers"))

    attempt_count = db.execute(
        "SELECT COUNT(*) AS count FROM exam_attempts WHERE paper_id = ? AND user_id = ?",
        (paper_id, session["user"]["id"]),
    ).fetchone()["count"]
    max_attempts = paper["max_attempts"] or 1
    if attempt_count >= max_attempts:
        existing_attempt = get_existing_attempt(paper_id, session["user"]["id"])
        if existing_attempt is not None:
            flash("这份试卷的作答次数已用完，系统已打开最近一次成绩。", "success")
            return redirect(url_for("exam_result", attempt_id=existing_attempt["id"]))

    if request.method == "POST":
        score = 0
        answer_rows = []
        for item in items:
            answer = submitted_answer_for(item)
            correct = is_answer_correct(item["question_type"], answer, item["answer"])
            item_score = item["score"] if correct else 0
            score += item_score
            answer_rows.append((item["paper_question_id"], answer, int(correct), item_score))

        cursor = db.execute(
            """
            INSERT INTO exam_attempts (paper_id, user_id, score, total_score)
            VALUES (?, ?, ?, ?)
            """,
            (paper_id, session["user"]["id"], score, paper["total_score"]),
        )
        attempt_id = cursor.lastrowid
        db.executemany(
            """
            INSERT INTO exam_answers (attempt_id, paper_question_id, answer, is_correct, score)
            VALUES (?, ?, ?, ?, ?)
            """,
            [(attempt_id, *row) for row in answer_rows],
        )
        db.commit()
        return redirect(url_for("exam_result", attempt_id=attempt_id))

    return render_template(
        "take_paper.html",
        paper=paper,
        items=items,
        attempt_number=attempt_count + 1,
    )


@app.get("/exams/<int:attempt_id>/result")
@login_required
def exam_result(attempt_id):
    db = get_db()
    if current_user_is_teacher():
        attempt = db.execute(
            """
            SELECT a.*, p.title, u.username AS account,
                   COALESCE(NULLIF(u.display_name, ''), u.username) AS username
            FROM exam_attempts a
            JOIN papers p ON p.id = a.paper_id
            JOIN users u ON u.id = a.user_id
            WHERE a.id = ?
            """,
            (attempt_id,),
        ).fetchone()
    else:
        attempt = db.execute(
            """
            SELECT a.*, p.title, u.username AS account,
                   COALESCE(NULLIF(u.display_name, ''), u.username) AS username
            FROM exam_attempts a
            JOIN papers p ON p.id = a.paper_id
            JOIN users u ON u.id = a.user_id
            WHERE a.id = ? AND a.user_id = ?
            """,
            (attempt_id, session["user"]["id"]),
        ).fetchone()
    if attempt is None:
        flash("成绩记录不存在。", "error")
        return redirect(url_for("dashboard"))
    details = db.execute(
        """
        SELECT ea.answer, ea.is_correct, ea.score, pq.sort_order, pq.score AS full_score,
               q.content, q.answer AS correct_answer, q.analysis, q.question_type
        FROM exam_answers ea
        JOIN paper_questions pq ON pq.id = ea.paper_question_id
        JOIN questions q ON q.id = pq.question_id
        WHERE ea.attempt_id = ? ORDER BY pq.sort_order
        """,
        (attempt_id,),
    ).fetchall()
    return render_template("exam_result.html", attempt=attempt, details=details)


def current_user_is_teacher():
    return session.get("user", {}).get("role") == "teacher"


def parse_optional_score(value, label):
    if not value:
        return None, None
    try:
        score = int(value)
    except ValueError:
        return None, f"{label}必须是数字。"
    if score < 0:
        return None, f"{label}不能小于 0。"
    return score, None


def build_result_stats(attempts):
    if not attempts:
        return {
            "attempt_count": 0,
            "student_count": 0,
            "avg_score": 0,
            "best_score": 0,
            "lowest_score": 0,
            "pass_rate": 0,
        }

    attempt_count = len(attempts)
    scores = [row["score"] for row in attempts]
    pass_count = sum(
        1
        for row in attempts
        if row["total_score"] and row["score"] * 100 / row["total_score"] >= 60
    )
    return {
        "attempt_count": attempt_count,
        "student_count": len({row["user_id"] for row in attempts}),
        "avg_score": round(sum(scores) / attempt_count, 1),
        "best_score": max(scores),
        "lowest_score": min(scores),
        "pass_rate": round(pass_count * 100 / attempt_count, 1),
    }


def build_score_buckets(attempts):
    buckets = [
        {"label": "90-100分", "min": 90, "max": 100, "count": 0},
        {"label": "80-89分", "min": 80, "max": 89, "count": 0},
        {"label": "70-79分", "min": 70, "max": 79, "count": 0},
        {"label": "60-69分", "min": 60, "max": 69, "count": 0},
        {"label": "0-59分", "min": 0, "max": 59, "count": 0},
    ]
    for row in attempts:
        total_score = row["total_score"] or 0
        if not total_score:
            continue
        score_rate = row["score"] * 100 / total_score
        for bucket in buckets:
            if bucket["min"] <= score_rate <= bucket["max"]:
                bucket["count"] += 1
                break
    return buckets


def build_score_curve(score_buckets):
    if not score_buckets:
        return {"points": "", "area_points": "", "peak": None}
    max_count = max([bucket["count"] for bucket in score_buckets] + [1])
    count = len(score_buckets)
    points = []
    peak = None
    for index, bucket in enumerate(score_buckets):
        x = 70 + index * (1090 / (count - 1 if count > 1 else 1))
        y = 240 - (bucket["count"] / max_count * 180)
        point = {"x": round(x, 2), "y": round(y, 2), "bucket": bucket}
        points.append(point)
        if peak is None or bucket["count"] > peak["bucket"]["count"]:
            peak = point
    point_text = " ".join(f"{point['x']},{point['y']}" for point in points)
    return {
        "points": point_text,
        "area_points": f"70,240 {point_text} 1160,240",
        "peak": peak,
    }


def replacement_choice_label(question):
    chapter = question["chapter"] or "未分类"
    text = question["content"].strip()
    if len(text) > 36:
        text = text[:36] + "..."
    return f"{text} · {chapter} · {DIFFICULTY_LABELS[question['difficulty']]}"


def replacement_choice_payload(question, is_current=False):
    return {
        "id": "" if is_current else question["id"],
        "content": question["content"],
        "chapter": question["chapter"] or "未分类",
        "difficulty": question["difficulty"],
        "difficulty_label": DIFFICULTY_LABELS[question["difficulty"]],
        "question_type": question["question_type"],
        "question_type_label": QUESTION_TYPE_LABELS[question["question_type"]],
        "label": "保持当前题目" if is_current else replacement_choice_label(question),
        "is_current": is_current,
    }


def build_replacement_choices(questions, subject, chapter):
    db = get_db()
    selected_ids = {question["id"] for question in questions}
    choices = []
    for question in questions:
        sql = """
            SELECT id, content, question_type, difficulty, chapter
            FROM questions
            WHERE subject = ? AND question_type = ?
        """
        params = [subject, question["question_type"]]
        if chapter:
            sql += " AND chapter = ?"
            params.append(chapter)
        sql += """
            ORDER BY
                CASE difficulty
                    WHEN 'easy' THEN 1
                    WHEN 'medium' THEN 2
                    WHEN 'hard' THEN 3
                    ELSE 9
                END,
                id DESC
        """
        rows = db.execute(sql, params).fetchall()
        row_choices = [replacement_choice_payload(question, is_current=True)]
        for row in rows:
            if row["id"] in selected_ids:
                continue
            row_choices.append(replacement_choice_payload(row))
        choices.append(row_choices)
    return choices


def apply_preview_replacements(questions, form, subject, chapter):
    db = get_db()
    final_questions = []
    used_ids = set()
    for index, question in enumerate(questions):
        raw_replacement = form.get(f"replace_{index}", "").strip()
        if not raw_replacement:
            chosen = question
        else:
            if not raw_replacement.isdigit():
                return "更换题目参数不正确，请重新预览。"
            replacement_id = int(raw_replacement)
            chosen = db.execute("SELECT * FROM questions WHERE id = ?", (replacement_id,)).fetchone()
            if chosen is None:
                return "更换题目不存在，请重新预览。"
            if chosen["subject"] != subject:
                return "更换题目必须保持同一科目，请重新预览。"
            if chapter and chosen["chapter"] != chapter:
                return "更换题目必须保持同一章节，请重新预览。"
            if chosen["question_type"] != question["question_type"]:
                return "更换题目必须保持同一题型，请重新预览。"
        if chosen["id"] in used_ids:
            return "更换后的题目不能重复，请重新预览。"
        used_ids.add(chosen["id"])
        final_questions.append(chosen)
    return final_questions


def validate_register_form(form):
    if not re.fullmatch(r"\d{8}", form["username"]):
        return "账号必须是 8 位数字。"
    if not valid_display_name(form["display_name"]):
        return "请填写 2-20 个字符的用户名。"
    if form["role"] not in {"student", "teacher"}:
        return "请选择有效的用户身份。"
    return None


def valid_display_name(display_name):
    return 2 <= len(display_name.strip()) <= 20


def get_chapters(subject=None):
    subject_filter = ""
    params = []
    if subject:
        subject_filter = "AND subject = ?"
        params.append(subject)
    rows = get_db().execute(
        f"""
        SELECT chapter FROM questions
        WHERE TRIM(chapter) != ''
        {subject_filter}
        GROUP BY chapter
        ORDER BY chapter
        """,
        params,
    ).fetchall()
    return [row["chapter"] for row in rows]


def get_subject_chapters():
    rows = get_db().execute(
        """
        SELECT subject, chapter FROM questions
        WHERE TRIM(chapter) != ''
        GROUP BY subject, chapter
        ORDER BY subject, chapter
        """
    ).fetchall()
    subject_chapters = {subject: [] for subject in SUBJECTS}
    for row in rows:
        subject_chapters.setdefault(row["subject"] or DEFAULT_SUBJECT, []).append(row["chapter"])
    return subject_chapters


def empty_availability_bucket():
    return {
        "total": 0,
        "by_type": {key: 0 for key in QUESTION_TYPE_LABELS},
        "by_difficulty": {key: 0 for key in DIFFICULTY_LABELS},
    }


def add_to_availability(bucket, question_type, difficulty):
    bucket["total"] += 1
    bucket["by_type"][question_type] = bucket["by_type"].get(question_type, 0) + 1
    bucket["by_difficulty"][difficulty] = bucket["by_difficulty"].get(difficulty, 0) + 1


def get_question_availability():
    rows = get_db().execute(
        """
        SELECT subject, chapter, question_type, difficulty
        FROM questions
        ORDER BY subject, chapter
        """
    ).fetchall()
    availability = {}
    for row in rows:
        subject = row["subject"] or DEFAULT_SUBJECT
        chapter = row["chapter"] or "未分类"
        availability.setdefault(subject, {"": empty_availability_bucket()})
        availability[subject].setdefault(chapter, empty_availability_bucket())
        add_to_availability(availability[subject][""], row["question_type"], row["difficulty"])
        add_to_availability(availability[subject][chapter], row["question_type"], row["difficulty"])
    for subject in SUBJECTS:
        availability.setdefault(subject, {"": empty_availability_bucket()})
    return availability


def build_paper_distribution(items):
    distribution = {
        "subjects": {},
        "chapters": {},
        "types": {},
        "difficulties": {},
    }
    for item in items:
        subject = item["subject"] or DEFAULT_SUBJECT
        chapter = item["chapter"] or "未分类"
        question_type = QUESTION_TYPE_LABELS[item["question_type"]]
        difficulty = DIFFICULTY_LABELS[item["difficulty"]]
        distribution["subjects"][subject] = distribution["subjects"].get(subject, 0) + 1
        distribution["chapters"][chapter] = distribution["chapters"].get(chapter, 0) + 1
        distribution["types"][question_type] = distribution["types"].get(question_type, 0) + 1
        distribution["difficulties"][difficulty] = distribution["difficulties"].get(difficulty, 0) + 1
    return distribution


def count_available_questions(subject, chapter=""):
    sql = "SELECT COUNT(*) AS count FROM questions WHERE subject = ?"
    params = [subject]
    if chapter:
        sql += " AND chapter = ?"
        params.append(chapter)
    return get_db().execute(sql, params).fetchone()["count"]


def analytics_group(field, user_filter, params, labels=None):
    rows = get_db().execute(
        f"""
        SELECT COALESCE(NULLIF({field}, ''), '未分类') AS label,
               COUNT(*) AS total,
               SUM(CASE WHEN ea.is_correct = 1 THEN 1 ELSE 0 END) AS correct
        FROM exam_answers ea
        JOIN exam_attempts a ON a.id = ea.attempt_id
        JOIN paper_questions pq ON pq.id = ea.paper_question_id
        JOIN questions q ON q.id = pq.question_id
        WHERE 1 = 1 {user_filter}
        GROUP BY COALESCE(NULLIF({field}, ''), '未分类')
        ORDER BY total DESC, label ASC
        """,
        params,
    ).fetchall()
    result = []
    for row in rows:
        label = labels.get(row["label"], row["label"]) if labels else row["label"]
        total = row["total"] or 0
        correct = row["correct"] or 0
        rate = round(correct * 100 / total, 1) if total else 0
        result.append({"label": label, "total": total, "correct": correct, "rate": rate})
    return result


def _radar_geometry(labels, values):
    if not labels:
        return {
            "labels": [],
            "values": [],
            "full_points": "",
            "value_points": "",
            "label_positions": [],
        }
    center_x = 120
    center_y = 110
    radius = 68
    full_points = []
    value_points = []
    label_positions = []
    count = len(labels)
    for index, (label, value) in enumerate(zip(labels, values)):
        angle = math.radians(-90 + index * (360 / count))
        full_x = center_x + radius * math.cos(angle)
        full_y = center_y + radius * math.sin(angle)
        full_points.append(f"{full_x:.1f},{full_y:.1f}")
        value_radius = max(0.0, min(100.0, value)) / 100.0 * radius
        value_x = center_x + value_radius * math.cos(angle)
        value_y = center_y + value_radius * math.sin(angle)
        value_points.append(f"{value_x:.1f},{value_y:.1f}")
        label_radius = radius + 24
        label_x = center_x + label_radius * math.cos(angle)
        label_y = center_y + label_radius * math.sin(angle)
        label_positions.append({"label": label, "x": round(label_x, 1), "y": round(label_y, 1)})
    return {
        "labels": labels,
        "values": values,
        "full_points": " ".join(full_points),
        "value_points": " ".join(value_points),
        "label_positions": label_positions,
    }


def build_ability_radar_data(user_id):
    rows = get_db().execute(
        """
        SELECT COALESCE(NULLIF(q.subject, ''), ?) AS subject,
               COALESCE(NULLIF(q.chapter, ''), '未分类') AS chapter,
               COUNT(*) AS total,
               SUM(CASE WHEN ea.is_correct = 1 THEN 1 ELSE 0 END) AS correct
        FROM exam_answers ea
        JOIN exam_attempts a ON a.id = ea.attempt_id
        JOIN paper_questions pq ON pq.id = ea.paper_question_id
        JOIN questions q ON q.id = pq.question_id
        WHERE a.user_id = ?
        GROUP BY q.subject, q.chapter
        """,
        (DEFAULT_SUBJECT, user_id),
    ).fetchall()
    subject_rates = {subject: {} for subject in SUBJECTS}
    for row in rows:
        total = row["total"] or 0
        correct = row["correct"] or 0
        rate = round(correct * 100 / total, 1) if total else 0
        subject_rates[row["subject"]][row["chapter"]] = rate

    datasets = []
    for subject in SUBJECTS:
        chapter_rates = subject_rates.get(subject, {})
        labels = list(chapter_rates.keys())
        values = [chapter_rates[label] for label in labels]
        geometry = _radar_geometry(labels, values)
        datasets.append({"subject": subject, **geometry, "has_data": bool(labels)})
    return datasets


def first_student_subject(user_id):
    row = get_db().execute(
        """
        SELECT q.subject
        FROM exam_answers ea
        JOIN exam_attempts a ON a.id = ea.attempt_id
        JOIN paper_questions pq ON pq.id = ea.paper_question_id
        JOIN questions q ON q.id = pq.question_id
        WHERE a.user_id = ?
        ORDER BY a.submitted_at ASC, a.id ASC
        LIMIT 1
        """,
        (user_id,),
    ).fetchone()
    return row["subject"] if row else DEFAULT_SUBJECT


def build_ability_trend(user_id, subject=None):
    student_class = get_db().execute(
        "SELECT class_id FROM users WHERE id = ?",
        (user_id,),
    ).fetchone()
    class_id = student_class["class_id"] if student_class else None
    sql = """
        SELECT DISTINCT a.id, a.score, a.total_score, a.submitted_at, a.paper_id, p.title,
               ROUND(a.score * 100.0 / NULLIF(a.total_score, 0), 1) AS rate
        FROM exam_attempts a
        JOIN papers p ON p.id = a.paper_id
        JOIN paper_questions pq ON pq.paper_id = a.paper_id
        JOIN questions q ON q.id = pq.question_id
        WHERE a.user_id = ?
    """
    params = [user_id]
    if subject:
        sql += " AND q.subject = ?"
        params.append(subject)
    sql += " ORDER BY a.submitted_at ASC, a.id ASC"
    rows = get_db().execute(sql, params).fetchall()
    points = []
    for index, row in enumerate(rows):
        class_avg = 0
        if class_id:
            avg_row = get_db().execute(
                """
                SELECT ROUND(AVG(score * 100.0 / NULLIF(total_score, 0)), 1) AS avg_rate
                FROM exam_attempts
                WHERE paper_id = ? AND user_id IN (
                    SELECT id FROM users WHERE class_id = ?
                )
                """,
                (row["paper_id"], class_id),
            ).fetchone()
            class_avg = avg_row["avg_rate"] or 0 if avg_row else 0
        points.append({
            "label": f"第{index + 1}次",
            "date": row["submitted_at"],
            "title": row["title"],
            "student_rate": row["rate"] or 0,
            "class_avg_rate": round(class_avg or 0, 1),
        })
    return points


def get_paper_items(paper_id):
    return get_db().execute(
        """
        SELECT pq.id AS paper_question_id, pq.sort_order, pq.score, q.*
        FROM paper_questions pq JOIN questions q ON q.id = pq.question_id
        WHERE pq.paper_id = ? ORDER BY pq.sort_order
        """,
        (paper_id,),
    ).fetchall()


def get_paper(paper_id):
    return get_db().execute(
        """
        SELECT p.*, COALESCE(NULLIF(u.display_name, ''), u.username) AS creator FROM papers p
        LEFT JOIN users u ON u.id = p.created_by WHERE p.id = ?
        """,
        (paper_id,),
    ).fetchone()


def get_paper_subject(paper_id):
    rows = get_db().execute(
        """
        SELECT DISTINCT COALESCE(NULLIF(q.subject, ''), ?) AS subject
        FROM paper_questions pq
        JOIN questions q ON q.id = pq.question_id
        WHERE pq.paper_id = ?
        ORDER BY subject
        """,
        (DEFAULT_SUBJECT, paper_id),
    ).fetchall()
    subjects = [row["subject"] for row in rows]
    return "、".join(subjects) if subjects else DEFAULT_SUBJECT


def get_existing_attempt(paper_id, user_id):
    return get_db().execute(
        """
        SELECT id FROM exam_attempts
        WHERE paper_id = ? AND user_id = ?
        ORDER BY submitted_at DESC, id DESC
        LIMIT 1
        """,
        (paper_id, user_id),
    ).fetchone()


def parse_options(options_text):
    return [line.strip() for line in (options_text or "").splitlines() if line.strip()]


def export_filename(paper, suffix):
    safe_title = re.sub(r"[\\/:*?\"<>|]+", "_", paper["title"]).strip() or "试卷"
    return f"{paper['id']}-{safe_title}.{suffix}"


def paper_subjects(items):
    subjects = sorted({item["subject"] or DEFAULT_SUBJECT for item in items})
    return "、".join(subjects)


def question_export_lines(item):
    lines = [
        f"第 {item['sort_order']} 题 [{QUESTION_TYPE_LABELS[item['question_type']]}] "
        f"{item['score']} 分 {DIFFICULTY_LABELS[item['difficulty']]}",
        item["content"],
    ]
    if item["question_type"] in {"single", "multiple"}:
        lines.extend(parse_options(item["options"]))
    elif item["question_type"] == "true_false":
        lines.extend(["对", "错"])
    lines.append(f"参考答案：{item['answer']}")
    if item["analysis"]:
        lines.append(f"解析：{item['analysis']}")
    return lines


def build_paper_docx(paper, items):
    from docx import Document
    from docx.shared import Pt
    from docx.oxml.ns import qn

    EXPORT_DIR.mkdir(exist_ok=True)
    path = EXPORT_DIR / export_filename(paper, "docx")
    document = Document()
    styles = document.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"].font.size = Pt(11)
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    title = document.add_heading(paper["title"], level=0)
    title.alignment = 1
    document.add_paragraph(
        f"科目：{paper_subjects(items)}    总分：{paper['total_score']} 分    "
        f"题量：{len(items)} 题    创建人：{paper['creator']}"
    )
    document.add_paragraph("说明：本导出版本供教师阅卷、打印或归档使用，包含参考答案与解析。")

    for item in items:
        document.add_paragraph("")
        for index, line in enumerate(question_export_lines(item)):
            paragraph = document.add_paragraph(line)
            if index == 0:
                for run in paragraph.runs:
                    run.bold = True
    document.save(path)
    return path


def chinese_pdf_font():
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    candidates = [
        Path("C:/Windows/Fonts/simhei.ttf"),
        Path("C:/Windows/Fonts/NotoSansSC-VF.ttf"),
        Path("C:/Windows/Fonts/NotoSerifSC-VF.ttf"),
    ]
    for font_path in candidates:
        if font_path.exists():
            try:
                pdfmetrics.registerFont(TTFont("PaperExportFont", str(font_path)))
                return "PaperExportFont"
            except Exception:
                continue
    return "Helvetica"


def paragraph_text(value):
    return escape(str(value)).replace("\n", "<br/>")


def build_paper_pdf(paper, items):
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    EXPORT_DIR.mkdir(exist_ok=True)
    path = EXPORT_DIR / export_filename(paper, "pdf")
    font_name = chinese_pdf_font()
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "PaperTitle",
        parent=styles["Title"],
        fontName=font_name,
        fontSize=18,
        leading=24,
        alignment=TA_CENTER,
        spaceAfter=12,
    )
    meta_style = ParagraphStyle(
        "PaperMeta",
        parent=styles["Normal"],
        fontName=font_name,
        fontSize=10,
        leading=15,
        spaceAfter=8,
    )
    question_style = ParagraphStyle(
        "PaperQuestion",
        parent=styles["Normal"],
        fontName=font_name,
        fontSize=10.5,
        leading=16,
        spaceBefore=7,
        spaceAfter=3,
    )
    strong_style = ParagraphStyle(
        "PaperStrong",
        parent=question_style,
        fontName=font_name,
        fontSize=10.5,
        leading=16,
        spaceBefore=10,
        spaceAfter=3,
    )

    doc = SimpleDocTemplate(
        str(path),
        pagesize=A4,
        rightMargin=18 * mm,
        leftMargin=18 * mm,
        topMargin=16 * mm,
        bottomMargin=16 * mm,
    )
    story = [
        Paragraph(paragraph_text(paper["title"]), title_style),
        Paragraph(
            paragraph_text(
                f"科目：{paper_subjects(items)}    总分：{paper['total_score']} 分    "
                f"题量：{len(items)} 题    创建人：{paper['creator']}"
            ),
            meta_style,
        ),
        Paragraph("说明：本导出版本供教师阅卷、打印或归档使用，包含参考答案与解析。", meta_style),
        Spacer(1, 4 * mm),
    ]
    for item in items:
        for index, line in enumerate(question_export_lines(item)):
            style = strong_style if index == 0 else question_style
            story.append(Paragraph(paragraph_text(line), style))
    doc.build(story)
    return path


def submitted_answer_for(item):
    field_name = f"question_{item['paper_question_id']}"
    if item["question_type"] == "multiple":
        return ",".join(request.form.getlist(field_name)).strip()
    return request.form.get(field_name, "").strip()


def normalize_answer(value):
    normalized = value.strip().replace("，", ",").replace(" ", "").upper()
    answer_aliases = {
        "对": "正确",
        "错": "错误",
        "TRUE": "正确",
        "FALSE": "错误",
        "T": "正确",
        "F": "错误",
    }
    return answer_aliases.get(normalized, normalized)


def is_answer_correct(question_type, submitted, expected):
    submitted = normalize_answer(submitted)
    expected = normalize_answer(expected)
    if question_type == "multiple":
        return set(filter(None, submitted.split(","))) == set(filter(None, expected.split(",")))
    return submitted == expected


def build_paper_from_form(form):
    db = get_db()
    try:
        counts = {
            "single": int(form.get("single_count", 0)),
            "multiple": int(form.get("multiple_count", 0)),
            "true_false": int(form.get("true_false_count", 0)),
            "fill": int(form.get("fill_count", 0)),
        }
        difficulty_counts = {
            "easy": int(form.get("easy_count", 0)),
            "medium": int(form.get("medium_count", 0)),
            "hard": int(form.get("hard_count", 0)),
        }
        score_per_question = int(form.get("score_per_question", 5))
        duration_minutes = int(form.get("duration_minutes", 30))
        max_attempts = int(form.get("max_attempts", 1))
    except ValueError:
        return "请输入有效的数字。"
    subject = form.get("subject", DEFAULT_SUBJECT).strip()
    chapter = form.get("chapter", "").strip()
    if min(*counts.values(), *difficulty_counts.values(), score_per_question, duration_minutes) < 0:
        return "数量不能为负数。"
    if duration_minutes <= 0:
        return "考试时长必须大于 0 分钟。"
    if max_attempts <= 0:
        return "限做次数必须大于 0。"
    if sum(counts.values()) == 0:
        return "请至少设置一种题型数量。"
    if not subject:
        return "请先选择组卷科目。"

    selected = []
    seen = set()
    expected_total = sum(counts.values()) + sum(difficulty_counts.values())
    available_total = count_available_questions(subject, chapter)
    if available_total < expected_total:
        scope = f"科目“{subject}”章节“{chapter}”中" if chapter else f"科目“{subject}”中"
        return f"{scope}总题量不足，需要 {expected_total} 道，当前只有 {available_total} 道。"

    for qtype, count in counts.items():
        if count:
            sql = "SELECT * FROM questions WHERE subject = ? AND question_type = ?"
            params = [subject, qtype]
            if chapter:
                sql += " AND chapter = ?"
                params.append(chapter)
            sql += " ORDER BY RANDOM()"
            rows = db.execute(
                sql,
                params,
            ).fetchall()
            available_rows = [row for row in rows if row["id"] not in seen]
            if len(available_rows) < count:
                scope = f"科目“{subject}”章节“{chapter}”中" if chapter else f"科目“{subject}”中"
                return f"{scope}可用不重复{QUESTION_TYPE_LABELS[qtype]}数量不足，当前只有 {len(available_rows)} 道。"
            for row in available_rows[:count]:
                seen.add(row["id"])
                selected.append(row)

    for difficulty, count in difficulty_counts.items():
        if count:
            sql = "SELECT * FROM questions WHERE subject = ? AND difficulty = ?"
            params = [subject, difficulty]
            if chapter:
                sql += " AND chapter = ?"
                params.append(chapter)
            sql += " ORDER BY RANDOM()"
            rows = db.execute(
                sql,
                params,
            ).fetchall()
            available_rows = [row for row in rows if row["id"] not in seen]
            if len(available_rows) < count:
                scope = f"科目“{subject}”章节“{chapter}”中" if chapter else f"科目“{subject}”中"
                return f"{scope}可用不重复{DIFFICULTY_LABELS[difficulty]}题数量不足，当前只有 {len(available_rows)} 道。"
            for row in available_rows[:count]:
                seen.add(row["id"])
                selected.append(row)

    if len(selected) != expected_total:
        return f"组卷题量异常，需要 {expected_total} 道，实际抽取 {len(selected)} 道，请调整规则后重试。"

    selected = sorted(
        selected,
        key=lambda row: (QUESTION_TYPE_RANK.get(row["question_type"], 99), row["id"]),
    )

    return {
        "questions": selected,
        "score_per_question": score_per_question,
        "duration_minutes": duration_minutes,
        "max_attempts": max_attempts,
        "total_score": len(selected) * score_per_question,
    }


def save_paper(user_id, title, result, class_ids=None):
    db = get_db()
    cursor = db.execute(
        "INSERT INTO papers (title, total_score, duration_minutes, max_attempts, created_by) VALUES (?, ?, ?, ?, ?)",
        (
            title,
            result["total_score"],
            result.get("duration_minutes", 30),
            result.get("max_attempts", 1),
            user_id,
        ),
    )
    paper_id = cursor.lastrowid
    for class_id in class_ids or []:
        db.execute(
            "INSERT INTO paper_classes (paper_id, class_id) VALUES (?, ?)",
            (paper_id, class_id),
        )
    db.executemany(
        """
        INSERT INTO paper_questions (paper_id, question_id, score, sort_order)
        VALUES (?, ?, ?, ?)
        """,
        [
            (paper_id, question["id"], result["score_per_question"], index)
            for index, question in enumerate(result["questions"], start=1)
        ],
    )
    db.commit()
    return paper_id


def empty_question_form():
    return {
        "content": "",
        "question_type": "single",
        "difficulty": "easy",
        "subject": DEFAULT_SUBJECT,
        "chapter": "",
        "options": "",
        "answer": "",
        "analysis": "",
    }


def question_form_data():
    return {
        "content": request.form.get("content", "").strip(),
        "question_type": request.form.get("question_type", "").strip(),
        "difficulty": request.form.get("difficulty", "").strip(),
        "subject": request.form.get("subject", DEFAULT_SUBJECT).strip(),
        "chapter": request.form.get("chapter", "").strip(),
        "options": request.form.get("options", "").strip(),
        "answer": request.form.get("answer", "").strip(),
        "analysis": request.form.get("analysis", "").strip(),
    }


def validate_question_form(form):
    if not form["content"]:
        return "题目内容不能为空。"
    if form["question_type"] not in QUESTION_TYPE_LABELS:
        return "请选择有效的题目类型。"
    if form["difficulty"] not in DIFFICULTY_LABELS:
        return "请选择有效的难度。"
    if not form["subject"]:
        return "请选择或填写科目。"
    if form["question_type"] in {"single", "multiple"} and not form["options"]:
        return "选择题需要填写选项。"
    if not form["answer"]:
        return "参考答案不能为空。"
    return None


if __name__ == "__main__":
    with app.app_context():
        init_db()
    app.run(host="0.0.0.0", port=5000, debug=True)
